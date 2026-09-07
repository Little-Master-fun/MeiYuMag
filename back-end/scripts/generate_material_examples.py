"""Regenerate clearly labelled fictional references, without changing original forms.

Requires python-docx and reportlab. Render and visually verify outputs after running.
"""
from pathlib import Path
import os

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from reportlab.lib.styles import ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import SimpleDocTemplate, Paragraph

OUTPUT = Path(__file__).resolve().parents[1] / "app" / "static" / "templates"
INTRO = "本文件为虚构填写示例，仅供理解材料内容，不是正式审批表，也不具有签章效力。请勿原样提交。正式申请时请替换所有示例信息，并按场地要求使用原表。"
EXAMPLES = {
    "meiyu_application_example": ("场地申请填写示例", [
        ("活动基本信息", ["活动名称：秋季阅读交流会（示例）", "申请场地：会议室", "申请组织：示例读书社", "联系人：示例联系人；联系电话：请填写真实号码", "使用时间：2030年10月20日 14:00 至 16:00（请改为实际日期）", "预计人数：20人"]),
        ("活动内容", ["围绕一本书开展分享与讨论。14:00签到，14:15主题分享，15:15自由交流，15:50整理场地，16:00结束。以上时间为示例。"]),
        ("使用与保障安排", ["拟使用桌椅与投影设备，具体需求以实际场地提供的设施为准。活动负责人负责现场秩序、人员疏散和结束后的清理。"]),
        ("提交前核对", ["请在正式申请表内填写组织、联系方式、准确场地名称及起止时间，确保与页面选择的场地和日期一致。初审上传DOCX，初审通过后再按要求递交真实签章材料。"]),
    ]),
    "yueyuan_plan_example": ("悦园三楼活动策划书示例", [
        ("活动基本信息", ["活动名称：秋日手作交流活动（示例）", "申请场地：悦园三楼", "主办组织：示例艺术社；负责人：示例联系人", "联系电话：请填写真实号码", "使用时间：2030年10月20日 14:00 至 16:00（请改为实际日期）", "预计人数：30人"]),
        ("活动目的与流程", ["以纸艺体验开展交流。14:00签到与安全说明，14:15作品制作，15:20交流展示，15:45整理桌面，16:00结束。"]),
        ("人员与设备安排", ["负责人统筹活动；现场工作人员分别负责签到、秩序与清洁。设备需求：桌椅、投影（需事先确认可用）。如实际使用电器，请列出设备名称、功率、数量与用电安排。"]),
        ("安全与应急安排", ["按实际人数安排活动区域，保持通道畅通；负责人提前熟悉出口，出现异常时停止活动并联系管理人员。不要用本示例代替实际风险评估。"]),
        ("材料准备", ["初审提交DOCX策划书。通过后按页面清单提交策划书、安全责任书、检查清单、用电承诺书及各自的真实签章扫描件。本示例不含签名或印章。"]),
    ]),
    "supporting_material_example": ("补充说明填写示例", [
        ("对应申请", ["申请编号：请填写原申请编号", "活动名称：秋季阅读交流会（示例）", "申请场地：会议室；申请组织：示例读书社", "活动时间：2030年10月20日 14:00 至 16:00（请改为实际日期）"]),
        ("本次补充内容", ["示例要求：请补充设备需求及现场人员安排。", "说明示例：活动拟使用投影一台，不自带大功率电器。安排两名工作人员负责签到、现场秩序与结束后的清洁；具体设备及人员请按实际情况修改。"]),
        ("关联附件", ["如需提供证明材料，请附上真实、清晰的原件或扫描件，并说明其与本次申请的关系。本说明不替代具有证明效力的原始材料。"]),
        ("提交方式", ["从原申请的补交入口进入信箱。上传后选择管理员要求的材料用途，不要把说明文件标记为签章扫描件；核对清单齐全后封缄送出。"]),
    ]),
    "key_borrow_example": ("钥匙借用填写示例", [
        ("借用信息", ["钥匙名称：会议室钥匙（示例，请填写实际钥匙名称）", "借用组织：示例读书社", "借用人：示例联系人；联系电话：请填写真实号码", "借用时间：2030年10月20日 13:30（请改为实际日期）", "预计归还时间：2030年10月20日 16:30（请改为实际日期）"]),
        ("借用事由", ["用于已安排的阅读交流活动，布置结束后开展活动，活动结束后整理场地并归还钥匙。借用事由须与实际使用安排一致。"]),
        ("准备与提交", ["将真实的借用信息填写在文档中并导出PDF，通过钥匙借用入口上传一份PDF。系统提取信息后由管理员审核，上传成功不代表借用已获批准。"]),
        ("重要说明", ["本文件仅用于展示信息组织方式，不是学校发布的正式借用表。正式办理格式、签字及归还要求以管理员通知为准。本示例不包含任何有效签章。"]),
    ]),
}


def make_docx(stem, title, sections):
    document = Document()
    page = document.sections[0]
    page.page_width, page.page_height = Inches(8.5), Inches(11)
    page.top_margin = page.bottom_margin = Inches(.72)
    page.left_margin = page.right_margin = Inches(.8)
    for name in ["Normal", "Title", "Heading 1"]:
        style = document.styles[name]
        style.font.name = "Arial Unicode MS"
        fonts = style.element.get_or_add_rPr().rFonts
        for key in list(fonts.attrib):
            del fonts.attrib[key]
        for key in ["ascii", "hAnsi", "eastAsia", "cs"]:
            fonts.set(qn(f"w:{key}"), "Arial Unicode MS")
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.size = Pt(11 if name == "Normal" else 22 if name == "Title" else 14)
        style.paragraph_format.space_after = Pt(7)
        style.paragraph_format.line_spacing = 1.2
        if stem == "yueyuan_plan_example":
            style.paragraph_format.space_before = Pt(10 if name == "Heading 1" else 0)
            style.paragraph_format.space_after = Pt(5)
            style.paragraph_format.line_spacing = 1.1
    document.add_paragraph(title, "Title")
    document.add_paragraph(INTRO)
    for heading, paragraphs in sections:
        document.add_heading(heading, level=1)
        for text in paragraphs:
            document.add_paragraph(text)
    document.core_properties.author = "美育系统"
    document.core_properties.title = title
    document.core_properties.subject = "虚构填写参考 请勿原样提交"
    # Remove inherited template borders, including the blue title rule.
    for root in [document.styles.element, document.element]:
        for border in root.iter(qn("w:pBdr")):
            border.getparent().remove(border)
    document.save(OUTPUT / f"{stem}.docx")


def make_key_pdf():
    # Embed a Chinese-capable subset, so PDF readers need no external CJK pack.
    pdfmetrics.registerFont(TTFont("ExampleChinese", os.environ.get("MEIYU_EXAMPLE_FONT", "/System/Library/Fonts/Supplemental/Arial Unicode.ttf")))
    body = ParagraphStyle("body", fontName="ExampleChinese", fontSize=11, leading=18, spaceAfter=8)
    heading = ParagraphStyle("heading", parent=body, fontSize=14, leading=20, spaceBefore=14, spaceAfter=8)
    title_style = ParagraphStyle("title", parent=body, fontSize=22, leading=30, spaceAfter=16)
    title, sections = EXAMPLES["key_borrow_example"]
    story = [Paragraph(title, title_style), Paragraph(INTRO, body)]
    for name, paragraphs in sections:
        story.append(Paragraph(name, heading))
        story.extend(Paragraph(text, body) for text in paragraphs)
    SimpleDocTemplate(str(OUTPUT / "key_borrow_example.pdf"), pagesize=(612, 792),
                      leftMargin=58, rightMargin=58, topMargin=52, bottomMargin=52,
                      title=title, author="美育系统").build(story)


if __name__ == "__main__":
    OUTPUT.mkdir(parents=True, exist_ok=True)
    for stem, (title, sections) in EXAMPLES.items():
        make_docx(stem, title, sections)
    make_key_pdf()
