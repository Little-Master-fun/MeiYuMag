from io import BytesIO

from docx import Document
from fastapi import HTTPException, UploadFile, status


class WordParserService:
    async def extract_text(self, file: UploadFile) -> str:
        if not file.filename or not file.filename.lower().endswith(".docx"):
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="Only .docx Word documents are supported for text extraction",
            )

        content = await file.read()
        await file.seek(0)
        if not content:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="Uploaded Word document is empty",
            )

        try:
            document = Document(BytesIO(content))
        except Exception as exc:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="Could not parse Word document",
            ) from exc

        parts: list[str] = []
        for paragraph in document.paragraphs:
            text = paragraph.text.strip()
            if text:
                parts.append(text)

        for table in document.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if cells:
                    parts.append(" | ".join(cells))

        parsed_text = "\n".join(parts).strip()
        if not parsed_text:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="No readable text found in Word document",
            )
        return parsed_text


word_parser_service = WordParserService()
